import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from timezonefinder import TimezoneFinder
from datetime import datetime
import pytz
import io
import os

st.set_page_config(page_title="Conversor de Planilhas IBUTG", layout="centered")
st.title("Conversor de Planilhas IBUTG (08h às 17h)")
st.write("Envie a planilha modelo e os arquivos do INMET para gerar uma nova planilha com os dados no horário local.")

modelo_file = st.file_uploader("📄 Envie a planilha modelo IBUTG (.xlsx)", type=["xlsx"], key="modelo")
csv_files = st.file_uploader("📑 Envie os arquivos do INMET (.csv, .xls, .xlsx)", type=["csv", "xls", "xlsx"], accept_multiple_files=True, key="inmet")

if modelo_file and csv_files:
    modelo_bytes = modelo_file.read()
    for uploaded_file in csv_files:
        st.markdown(f"---\n📂 Processando arquivo: `{uploaded_file.name}`")
        try:
            if uploaded_file.name.lower().endswith('.csv'):
                content = uploaded_file.read().decode("latin1")
                lines = content.splitlines()
                latitude = longitude = None
                for line in lines[:10]:
                    if 'LATITUDE' in line.upper():
                        latitude = float(line.split(':')[-1].replace(';', '').replace(',', '.'))
                    if 'LONGITUDE' in line.upper():
                        longitude = float(line.split(':')[-1].replace(';', '').replace(',', '.'))
                if latitude is None or longitude is None:
                    st.error(f"❌ Latitude ou longitude não encontrada no cabeçalho de `{uploaded_file.name}`.")
                    continue
                tf = TimezoneFinder()
                tz = pytz.timezone(tf.timezone_at(lng=longitude, lat=latitude))
                df = pd.read_csv(io.StringIO(content), sep=';', skiprows=8, encoding='latin1')
            else:
                st.warning(f"⚠️ Arquivo `{uploaded_file.name}` ignorado: apenas arquivos CSV são suportados nesta versão.")
                continue

            df = df[[
                'DATA (YYYY-MM-DD)', 'HORA (UTC)',
                'TEMPERATURA DO AR - BULBO SECO, HORARIA (°C)',
                'TEMPERATURA DO PONTO DE ORVALHO (°C)',
                'UMIDADE RELATIVA DO AR, HORARIA (%)',
                'VENTO, VELOCIDADE HORARIA (m/s)'
            ]]
            for col in df.columns[2:]:
                df[col] = pd.to_numeric(df[col].astype(str).str.replace(',', '.'), errors='coerce')

            df['data_hora_utc'] = pd.to_datetime(df['DATA (YYYY-MM-DD)'] + ' ' + df['HORA (UTC)'], format='%Y-%m-%d %H:%M', utc=True)
            df['data_hora_local'] = df['data_hora_utc'].dt.tz_convert(tz)
            df['hora_local'] = df['data_hora_local'].dt.hour

            df_filtrado = df[(df['hora_local'] >= 8) & (df['hora_local'] <= 17)].copy().sort_values(by='data_hora_local').reset_index(drop=True)

            dados_final = pd.DataFrame({
                'DATA': df_filtrado['data_hora_local'].dt.date,
                'HORA': df_filtrado['data_hora_local'].dt.strftime('%H:%M'),
                'TAR': df_filtrado['TEMPERATURA DO AR - BULBO SECO, HORARIA (°C)'],
                'TPO': df_filtrado['TEMPERATURA DO PONTO DE ORVALHO (°C)'],
                'UR': df_filtrado['UMIDADE RELATIVA DO AR, HORARIA (%)'],
                'VENTO': df_filtrado['VENTO, VELOCIDADE HORARIA (m/s)']
            })

            wb = load_workbook(io.BytesIO(modelo_bytes))
            ws = wb.active

            linha = 4
            for _, row in dados_final.iterrows():
                ibutg_val = ws[f'O{linha}'].value
                if ibutg_val is not None:
                    try:
                        ibutg_val = float(str(ibutg_val).replace(',', '.'))
                    except:
                        ibutg_val = 0
                else:
                    ibutg_val = 0
                if ibutg_val >= 0:
                    ws[f'D{linha}'] = row['DATA']
                    ws[f'E{linha}'] = row['HORA']
                    ws[f'G{linha}'] = row['TAR']
                    ws[f'H{linha}'] = row['TPO']
                    ws[f'I{linha}'] = row['UR']
                    ws[f'J{linha}'] = row['VENTO']
                    linha += 1

            nome_saida = f"IBUTG_{uploaded_file.name.replace('.CSV', '').replace(' ', '_')}.xlsx"
            wb.save(nome_saida)
            with open(nome_saida, 'rb') as f:
                st.download_button("⬇️ Baixar resultado", f, file_name=nome_saida)
        except Exception as e:
            st.error(f"Erro ao processar `{uploaded_file.name}`: {e}")
